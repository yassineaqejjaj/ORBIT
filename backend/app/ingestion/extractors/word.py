"""DOCX extraction with python-docx: headings → Markdown ``#``, lists → ``-``, tables → Markdown rows.

Paragraphs and tables are read in document order from the body XML.
"""

from __future__ import annotations

import io
import re
import zipfile

from docx import Document as load_docx
from docx.document import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.extractors import DOCX, ExtractedDocument, ExtractionError

_HEADING_STYLE = re.compile(r"^(?:heading|titre|überschrift|título)\s*(\d)", re.IGNORECASE)
_LIST_STYLE = re.compile(r"list|liste|puce|bullet|number", re.IGNORECASE)


def _heading_level(paragraph: Paragraph) -> int | None:
    style = paragraph.style
    name = (style.name if style is not None else "") or ""
    if name.lower() in {"title", "titre"}:
        return 1
    match = _HEADING_STYLE.match(name)
    if match:
        return max(1, min(6, int(match.group(1))))
    # Outline level set directly on the paragraph (documents without heading styles).
    ppr = paragraph._p.pPr
    if ppr is not None:
        outline = ppr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl")
        if outline is not None:
            value = outline.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
            if value is not None and value.isdigit() and int(value) < 6:
                return int(value) + 1
    return None


def _is_list(paragraph: Paragraph) -> bool:
    style = paragraph.style
    name = (style.name if style is not None else "") or ""
    if _LIST_STYLE.search(name):
        return True
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None


def _paragraph_markdown(paragraph: Paragraph) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    level = _heading_level(paragraph)
    if level is not None:
        return f"{'#' * level} {text}"
    if _is_list(paragraph):
        return f"- {text}"
    return text


def _cell_text(text: str) -> str:
    return " ".join(text.split()).replace("|", "/")


def _table_markdown(table: Table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        cells: list[str] = []
        previous = None
        for cell in row.cells:
            # Merged cells are repeated by python-docx: keep one copy.
            if previous is not None and cell._tc is previous:
                continue
            previous = cell._tc
            cells.append(_cell_text(cell.text))
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    lines = []
    for index, row in enumerate(rows):
        padded = row + [""] * (width - len(row))
        lines.append("| " + " | ".join(padded) + " |")
        if index == 0:
            lines.append("|" + "|".join(["---"] * width) + "|")
    return "\n".join(lines)


def _iter_blocks(document: DocxDocument) -> list[Paragraph | Table]:
    blocks: list[Paragraph | Table] = []
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            blocks.append(Paragraph(child, document))
        elif tag == "tbl":
            blocks.append(Table(child, document))
    return blocks


def extract_docx(data: bytes, filename: str | None = None) -> ExtractedDocument:
    try:
        document = load_docx(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ExtractionError(f"Fichier DOCX illisible ou corrompu : {exc}") from exc

    parts: list[str] = []
    for block in _iter_blocks(document):
        rendered = _paragraph_markdown(block) if isinstance(block, Paragraph) else _table_markdown(block)
        if rendered:
            parts.append(rendered)

    # Consecutive list items stay together; other blocks are separated by a blank line.
    lines: list[str] = []
    for part in parts:
        if lines and not (part.startswith("- ") and lines[-1].startswith("- ")):
            lines.append("")
        lines.append(part)

    core = document.core_properties
    title = (core.title or "").strip() or None
    author = (core.author or "").strip() or None
    return ExtractedDocument(
        text="\n".join(lines),
        format="docx",
        mime_type=DOCX,
        title=title,
        author=author,
        metadata={"paragraph_count": len(document.paragraphs), "table_count": len(document.tables)},
    )
