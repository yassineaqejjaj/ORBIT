"""Chantier F — sources (docs/AI_CONTEXT_ENGINEERING.md §F). Fictitious data only; no network (fake
transcription server through MockTransport, fake MCP servers)."""

from __future__ import annotations

import io
from datetime import date
from typing import Any

import httpx
import pytest
from docx import Document as new_docx

from app.config import settings
from app.db import get_sessionmaker
from app.enums import JobKind, MemoryKind
from app.ingestion import meetings, transcription
from app.ingestion.extractors import ExtractionError
from app.ingestion.queue import claim_next_job
from app.llm import guardrail
from app.memory import meeting_extraction as me
from app.seed import manifest as mf
from app.worker import Worker

API = "/api/v1/projects"
JSON = dict[str, Any]

VTT = """WEBVTT

NOTE fictif

00:00:01.000 --> 00:00:04.000
<v Claire Dubois>Bonjour à tous, point sur Atlas.</v>

00:00:04.500 --> 00:00:06.000
<v Claire Dubois>On commence.</v>

00:01:10.000 --> 00:01:20.000
<v Karim Benali>Décision : nous retenons OpenSearch pour la recherche hybride.</v>

00:03:00.000 --> 00:03:10.000
<v Karim Benali>Je m'en charge : je vais envoyer le chiffrage d'ici vendredi.</v>

00:04:00.000 --> 00:04:10.000
<v Claire Dubois>Action pour Lucas : préparer la démo avant le 20/10.</v>
"""

SRT = """1
00:00:01,000 --> 00:00:02,500
Alice Roux: Salut à tous.

2
00:01:02,000 --> 00:01:05,000
[Bruno Petit] Décision : le pilote démarre à Lyon.
"""


async def run_jobs(*kinds: JobKind) -> int:
    worker = Worker(concurrency=1, worker_id="test-sources")
    processed = 0
    for _ in range(50):
        async with get_sessionmaker()() as session:
            job = await claim_next_job(session, worker.worker_id, kinds=kinds)
        if job is None:
            break
        await worker._process(job.id)
        processed += 1
    return processed


# --- F1 parsers ------------------------------------------------------------------------------------------


def test_vtt_voices_merge_and_timestamps() -> None:
    transcript = meetings.parse(VTT.encode(), meetings.VTT)
    assert transcript.format == "vtt"
    assert transcript.speakers == ["Claire Dubois", "Karim Benali"]
    first = transcript.turns[0]
    assert (first.speaker, first.start, first.end) == ("Claire Dubois", 1.0, 6.0)
    assert first.text == "Bonjour à tous, point sur Atlas. On commence."
    assert transcript.duration == 250.0


def test_srt_name_prefix_and_brackets() -> None:
    transcript = meetings.parse(SRT.encode(), None, "copil.srt")
    assert [(t.speaker, t.start) for t in transcript.turns] == [("Alice Roux", 1.0), ("Bruno Petit", 62.0)]
    assert transcript.turns[1].text == "Décision : le pilote démarre à Lyon."


def test_plain_text_layouts() -> None:
    text = (
        "Date : 2026-10-01\n"
        "Alice Martin: Bonjour\n"
        "Décision : on part sur la V2.\n"  # label continuation, not a speaker
        "[00:05:00] Chloé Petit: Point suivant.\n"
        "00:06:00\n"  # Google Meet standalone timestamp
        "David Roy: Je confirme.\n"
        "Eve Durand   0:07\n"  # new Teams header line
        "Texte de Eve.\n"
        "0:00:08.120 --> 0:00:09.450\n"  # classic Teams cue, speaker, text
        "Farid Haddad\n"
        "Texte de Farid.\n"
    )
    turns = meetings.parse(text.encode(), "text/plain").turns
    assert turns[0].speaker is None and turns[0].text.startswith("Date")
    assert turns[1].speaker == "Alice Martin" and "Décision : on part sur la V2." in turns[1].text.split("\n")
    assert [(t.speaker, t.start) for t in turns[2:]] == [
        ("Chloé Petit", 300.0),
        ("David Roy", 360.0),
        ("Eve Durand", 7.0),
        ("Farid Haddad", 8.12),
    ]


def test_docx_teams_layout() -> None:
    document = new_docx()
    document.add_paragraph("Réunion Atlas")
    document.add_paragraph("Alice Martin   0:03\nBonjour, on démarre.")
    document.add_paragraph("Bruno Petit   1:15")
    document.add_paragraph("Décision : on garde le QR code.")
    buffer = io.BytesIO()
    document.save(buffer)
    transcript = meetings.parse(buffer.getvalue(), None, "transcription.docx")
    assert transcript.format == "docx"
    assert [(t.speaker, t.start) for t in transcript.turns if t.speaker] == [
        ("Alice Martin", 3.0),
        ("Bruno Petit", 75.0),
    ]


def test_render_keeps_speakers_and_timestamps_and_metadata() -> None:
    extracted = meetings.extract_meeting(
        VTT.encode(), meetings.VTT, title="Comité", date="2026-10-07", participants=["Lucas Morel"]
    )
    assert extracted.format == "transcript"
    assert "Participants : Lucas Morel, Claire Dubois, Karim Benali" in extracted.text
    assert "[00:01:10] Karim Benali : Décision : nous retenons OpenSearch" in extracted.text
    meta = extracted.metadata["meeting"]
    # Consecutive cues of one speaker are merged into one turn (short silence only).
    assert meta["speakers"] == ["Claire Dubois", "Karim Benali"] and meta["turn_count"] == 4
    assert meta["turns"][1] == {
        "speaker": "Karim Benali",
        "text": "Décision : nous retenons OpenSearch pour la recherche hybride.",
        "start": 70.0,
        "end": 80.0,
    }
    with pytest.raises(ExtractionError):
        meetings.parse(b"WEBVTT\n\n", meetings.VTT)


# --- F1 attribution ----------------------------------------------------------------------------------------


def test_meeting_statements_attribution() -> None:
    transcript = meetings.parse(VTT.encode(), meetings.VTT)
    text = meetings.render(transcript, title="Comité", date="2026-10-07")
    statements = me.extract_meeting_statements(text, transcript.speakers, me.TurnState(), date(2026, 10, 7))
    decisions = [s for s in statements if s.kind == MemoryKind.decision]
    actions = [s for s in statements if s.kind == MemoryKind.action]
    assert (
        len(decisions) == 1 and decisions[0].decided_by == "Karim Benali" and decisions[0].explicit_decision
    )
    assert decisions[0].quote.startswith("[00:01:10] Karim Benali : ")
    assert len(actions) == 2
    karim, lucas = actions
    assert karim.action_meta == {
        "owner": "Karim Benali",
        "due_date": "2026-10-09",  # « vendredi » after Wednesday 7 October
        "due_text": "vendredi",
        "speaker": "Karim Benali",
        "timestamp": "00:03:00",
    }
    assert (
        karim.content == "Action — Karim Benali : Envoyer le chiffrage d'ici vendredi (échéance : 2026-10-09)"
    )
    assert lucas.action_meta["owner"] == "Lucas" and lucas.action_meta["due_date"] == "2026-10-20"
    assert lucas.action_meta["speaker"] == "Claire Dubois"


def test_action_detection_and_due_dates() -> None:
    speakers = ["Bruno Leroy", "Alice Martin"]
    ref = date(2026, 10, 7)
    assert me.detect_action("Bruno se charge de relancer le fournisseur.", "Alice Martin", speakers, ref) == (
        "relancer le fournisseur",
        {"owner": "Bruno Leroy", "due_date": None, "due_text": None},
    )
    assert me.detect_action("Je vais regarder ça.", "Alice Martin", speakers, ref) is None
    assert me.detect_action("On se charge de tout.", "Alice Martin", speakers, ref) is None
    assert me.resolve_due("à livrer au plus tard le 3 janvier", ref) == ("2027-01-03", "3 janvier")
    assert me.resolve_due("pour 2 personnes", ref) == (None, None)


# --- F1 transcription guardrail and size limit -------------------------------------------------------------


class FakeTranscriptionServer:
    def __init__(self) -> None:
        self.calls: list[httpx.Request] = []

    def __call__(self, request: httpx.Request) -> httpx.Response:
        self.calls.append(request)
        assert request.url.path == "/v1/audio/transcriptions"
        return httpx.Response(
            200,
            json={
                "text": "Bonjour. Décision : on garde Lyon.",
                "segments": [
                    {"start": 0.0, "end": 2.0, "text": " Bonjour."},
                    {"start": 2.0, "end": 5.5, "text": " Décision : on garde Lyon.", "speaker": "Alice"},
                ],
            },
        )


@pytest.fixture
def transcription_on(monkeypatch: pytest.MonkeyPatch) -> FakeTranscriptionServer:
    monkeypatch.setattr(settings, "transcription_base_url", "https://transcription.example/v1")
    monkeypatch.setattr(settings, "transcription_api_key", "test-key")
    monkeypatch.setattr(settings, "transcription_local", False)
    monkeypatch.setattr(settings, "llm_max_classification", 1)
    monkeypatch.setattr(settings, "transcription_max_mb", 1)
    return FakeTranscriptionServer()


async def test_transcription_calls_server_with_segments(transcription_on: FakeTranscriptionServer) -> None:
    transcript = await transcription.transcribe(
        b"RIFFfake",
        "point.wav",
        "audio/wav",
        classification=1,
        transport=httpx.MockTransport(transcription_on),
    )
    assert [(t.speaker, t.start, t.text) for t in transcript.turns] == [
        (None, 0.0, "Bonjour."),
        ("Alice", 2.0, "Décision : on garde Lyon."),
    ]
    [request] = transcription_on.calls
    assert request.headers["authorization"] == "Bearer test-key"
    assert b'name="response_format"' in request.content and b"verbose_json" in request.content


async def test_transcription_guardrail_and_size_limit(
    transcription_on: FakeTranscriptionServer, monkeypatch: pytest.MonkeyPatch
) -> None:
    transport = httpx.MockTransport(transcription_on)
    before = guardrail.skip_count()
    for level in (2, 3):
        with pytest.raises(transcription.TranscriptionError, match="garde-fou"):
            await transcription.transcribe(
                b"x", "a.mp3", "audio/mpeg", classification=level, transport=transport
            )
    assert guardrail.skip_count() == before + 2
    with pytest.raises(transcription.TranscriptionError, match="trop volumineux"):
        await transcription.transcribe(
            b"x" * (1024 * 1024 + 1), "a.mp3", "audio/mpeg", classification=0, transport=transport
        )
    assert transcription_on.calls == []  # nothing left ORBIT
    monkeypatch.setattr(settings, "transcription_local", True)  # self-hosted server: C3 allowed
    transcript = await transcription.transcribe(
        b"x", "a.mp3", "audio/mpeg", classification=3, transport=transport
    )
    assert transcript.turns and len(transcription_on.calls) == 1
    monkeypatch.setattr(settings, "transcription_base_url", "")
    with pytest.raises(transcription.TranscriptionError, match="non configurée"):
        transcription.check(10, 0)


# --- F1 end to end: import → pipeline → memory ---------------------------------------------------------------


async def test_meeting_import_pipeline_and_memory(
    admin_client: httpx.AsyncClient, project: JSON, monkeypatch: pytest.MonkeyPatch
) -> None:
    slug = project["slug"]
    response = await admin_client.post(
        f"{API}/{slug}/documents/meeting",
        files={"file": ("comite.vtt", VTT.encode(), "text/vtt")},
        data={"title": "Comité Atlas", "meeting_date": "2026-10-07", "participants": "Lucas Morel"},
    )
    assert response.status_code == 201, response.text
    doc_id = response.json()["id"]
    await run_jobs(JobKind.ingest)
    detail = (await admin_client.get(f"{API}/{slug}/documents/{doc_id}")).json()
    assert detail["status"] == "indexed", detail
    meeting = detail["metadata"]["meeting"]
    assert meeting["date"] == "2026-10-07" and meeting["participants"] == ["Lucas Morel"]
    assert meeting["speakers"] == ["Claire Dubois", "Karim Benali"] and len(meeting["turns"]) == 4
    texts = " ".join(c["text"] for c in detail["chunks"])
    assert "[00:01:10] Karim Benali : " in texts

    await run_jobs(JobKind.extract_memory)
    items = (await admin_client.get(f"{API}/{slug}/memory", params={"limit": 100})).json()["items"]
    decision = next(i for i in items if i["kind"] == "decision")
    assert decision["decided_by"] == "Karim Benali" and decision["status"] == "validated"
    actions = {i["action_meta"]["owner"]: i for i in items if i["kind"] == "action"}
    assert set(actions) == {"Karim Benali", "Lucas"}
    assert actions["Karim Benali"]["action_meta"]["due_date"] == "2026-10-09"

    # Audio: refused without a transcription server, and C2 audio never leaves ORBIT.
    audio = {"file": ("point.mp3", b"ID3fake", "audio/mpeg")}
    refused = await admin_client.post(f"{API}/{slug}/documents/meeting", files=audio)
    assert refused.status_code == 422 and "non configurée" in refused.text
    monkeypatch.setattr(settings, "transcription_base_url", "https://transcription.example/v1")
    monkeypatch.setattr(settings, "llm_max_classification", 1)
    monkeypatch.setattr(settings, "transcription_local", False)
    c2 = await admin_client.post(f"{API}/{slug}/documents/meeting", files=audio, data={"classification": "2"})
    assert c2.status_code == 422 and "garde-fou" in c2.text
    bad = await admin_client.post(
        f"{API}/{slug}/documents/meeting", files={"file": ("x.pdf", b"%PDF-1.4", "application/pdf")}
    )
    assert bad.status_code == 422


def test_seed_meeting_transcript() -> None:
    manifest = mf.load_manifest()
    assert mf.validate_manifest(manifest) == []
    [doc] = [d for d in manifest["documents"] if d["mode"] == "meeting"]
    name, content, content_type = mf.meeting_file(doc, _now())
    assert name.endswith(".vtt") and content_type == "text/vtt" and b"{{" not in content
    transcript = meetings.parse(content, meetings.VTT)
    assert set(transcript.speakers) == set(doc["participants"])
    text = meetings.render(transcript, title=doc["title"])
    statements = me.extract_meeting_statements(text, transcript.speakers, me.TurnState(), date(2026, 10, 7))
    decisions = [(s.decided_by, s.explicit_decision) for s in statements if s.kind == MemoryKind.decision]
    actions = [s.action_meta for s in statements if s.kind == MemoryKind.action]
    assert decisions == [("Inès Moreau", True), ("Camille Martin", True)]
    assert (
        len(actions) == 1 and actions[0]["owner"] == "Léo Bernard" and actions[0]["due_date"] == "2026-10-09"
    )


def _now() -> Any:
    from datetime import UTC, datetime

    return datetime(2026, 10, 8, 9, tzinfo=UTC)
